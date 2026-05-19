import io
import os
import json
import requests
import xml.etree.ElementTree as ET
from urllib.parse import urljoin, urlparse
import fitz
import docx
import pandas as pd
from pptx import Presentation
from bs4 import BeautifulSoup
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings, ChatGoogleGenerativeAI
from app.Shared.Supabase.Domain.Supabase import Supabase
from app.components.Ingestion.Domain.IngestionModels import (
    IngestionResponse,
    WebsiteIngestionRequest,
    WebsiteSubroutesRequest,
    WebsiteSubroutesResponse,
    IngestedDocResponse
)

class IngestionService:
    def __init__(self):
        self.supabase = Supabase()

    def scrape_website(self, url: str) -> str:
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
        }
        response = requests.get(url, headers=headers, timeout=15)
        soup = BeautifulSoup(response.text, 'html.parser')
        for script_or_style in soup(["script", "style"]):
            script_or_style.decompose()
        return soup.get_text(separator=' ', strip=True)

    def discover_sub_routes(self, root_url: str) -> list[str]:
        parsed_root = urlparse(root_url)
        domain = parsed_root.netloc
        scheme = parsed_root.scheme or 'https'
        base_url = root_url.rstrip('/') + '/'
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
        }
        
        sitemap_candidates = [
            f"{scheme}://{domain}/robots.txt",
            f"{scheme}://{domain}/sitemap.xml",
            f"{scheme}://{domain}/sitemap_index.xml",
            f"{scheme}://{domain}/wp-sitemap.xml"
        ]
        
        all_sitemap_pages = []
        for candidate in sitemap_candidates:
            try:
                resp = requests.get(candidate, headers=headers, timeout=10)
                if resp.status_code == 200:
                    if candidate.endswith('robots.txt'):
                        for line in resp.text.splitlines():
                            if line.strip().lower().startswith('sitemap:'):
                                sitemap_url = line.split(':', 1)[1].strip()
                                all_sitemap_pages.extend(self._parse_sitemap(sitemap_url, headers))
                    else:
                        all_sitemap_pages.extend(self._parse_sitemap(candidate, headers))
                if all_sitemap_pages:
                    break
            except Exception:
                continue

        if all_sitemap_pages:
            root_path = parsed_root.path.rstrip('/')
            filtered = []
            seen = set()
            for page_url in all_sitemap_pages:
                clean = page_url.split('#')[0].split('?')[0].rstrip('/')
                parsed_page = urlparse(clean)
                if parsed_page.netloc == domain and parsed_page.path.rstrip('/').startswith(root_path):
                    if clean == base_url.rstrip('/'):
                        continue
                    if clean not in seen:
                        seen.add(clean)
                        filtered.append(clean)
            return filtered

        all_sub_links = []
        seen = set()
        try:
            resp = requests.get(base_url, headers=headers, timeout=15)
            if resp.status_code == 200:
                soup = BeautifulSoup(resp.text, 'html.parser')
                for link in soup.find_all('a', href=True):
                    href = link['href']
                    full_url = urljoin(base_url, href).split('#')[0].split('?')[0]
                    parsed_url = urlparse(full_url)
                    if parsed_url.netloc == domain:
                        clean = full_url.rstrip('/')
                        if '/page/' in clean:
                            continue
                        if clean == base_url.rstrip('/'):
                            continue
                        if clean not in seen:
                            seen.add(clean)
                            all_sub_links.append(clean)
        except Exception:
            pass

        return all_sub_links

    def _parse_sitemap(self, sitemap_url: str, headers: dict, depth=0) -> list[str]:
        if depth > 3:
            return []
        try:
            resp = requests.get(sitemap_url, headers=headers, timeout=15)
            if resp.status_code != 200:
                return []
            root_el = ET.fromstring(resp.content)
            ns = {'sm': 'http://www.sitemaps.org/schemas/sitemap/0.9'}
            urls = []
            child_sitemaps = root_el.findall('.//sm:sitemap/sm:loc', ns)
            if child_sitemaps:
                for loc in child_sitemaps:
                    urls.extend(self._parse_sitemap(loc.text.strip(), headers, depth + 1))
            else:
                url_entries = root_el.findall('.//sm:url/sm:loc', ns)
                for loc in url_entries:
                    urls.append(loc.text.strip())
            return urls
        except Exception:
            return []

    def classify_document(self, text: str, source: str) -> dict:
        snippet = text[:2000]
        prompt = f"""Analyze the following document text and return a JSON object with these fields:
- "title": A short, descriptive title for this document (max 10 words)
- "category": One of: case-study, one-pager, placemat, report, blog, email, presentation, general
- "industry": One of: technology, healthcare, finance, insurance, legal, manufacturing, energy, retail, logistics, staffing, general
- "topics": A list of 2-5 specific topic keywords (e.g. ["ERP", "cloud migration", "SAP"])

Source/filename hint: {source}

Document text:
---
{snippet}
---

Respond with ONLY valid JSON, no markdown fences or extra text."""
        try:
            model = ChatGoogleGenerativeAI(model="gemini-2.5-flash-lite", temperature=0)
            response = model.invoke(prompt)
            result_text = response.content.strip().strip('`').replace('json\n', '')
            return json.loads(result_text)
        except Exception:
            return {
                "title": source,
                "category": "general",
                "industry": "general",
                "topics": []
            }

    def generate_summary(self, text: str, source: str) -> str:
        snippet = text[:3000]
        prompt = f"""Summarize the following document in exactly 1-2 sentences. Be specific about what it covers.
Do NOT use meta-language or passive voice. Specifically, do NOT start the summary with "This document...", "This case study...", "This article...", "A document about...", or similar.
Start directly with the core action or subject in an active voice (e.g., "Streamlines ERP deployment..." instead of "This document describes how to streamline...").

Source: {source}

Document text:
---
{snippet}
---

Respond with ONLY the summary text, nothing else."""
        try:
            model = ChatGoogleGenerativeAI(model="gemini-2.5-flash-lite", temperature=0)
            response = model.invoke(prompt)
            return response.content.strip()
        except Exception:
            return f"Document from {source}"

    def upload_file_to_storage(self, file_bytes: bytes, filename: str) -> str:
        storage_path = f"uploads/{filename}"
        ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else 'bin'
        content_types = {
            'pdf': 'application/pdf',
            'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
            'txt': 'text/plain',
            'md': 'text/markdown',
        }
        content_type = content_types.get(ext, 'application/octet-stream')
        self.supabase.client.storage.from_("documents").upload(
            path=storage_path,
            file=file_bytes,
            file_options={"content-type": content_type, "upsert": "true"}
        )
        return storage_path

    def save_manifest_entry(self, source: str, source_type: str, title: str, category: str, industry: str, summary: str, chunk_count: int):
        data = {
            "source": source,
            "source_type": source_type,
            "title": title,
            "category": category,
            "industry": industry,
            "summary": summary,
            "chunk_count": chunk_count
        }
        self.supabase.client.table("knowledge_manifest").upsert(data, on_conflict="source").execute()

    def process_and_store(self, text: str, metadata: dict) -> tuple[int, list[str]]:
        logs = []
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
        chunks = text_splitter.split_text(text)
        
        embeddings_model = GoogleGenerativeAIEmbeddings(
            model="models/gemini-embedding-2",
            task_type="retrieval_document",
            output_dimensionality=1536
        )
        
        embeddings = []
        for i, chunk in enumerate(chunks):
            try:
                emb = embeddings_model.embed_query(chunk)
                embeddings.append(emb)
            except Exception as e:
                logs.append(f"Error embedding chunk {i+1}: {str(e)}")

        records = []
        for chunk, emb in zip(chunks, embeddings):
            records.append({
                "content": chunk,
                "metadata": metadata,
                "embedding": emb
            })

        if records:
            result = self.supabase.client.table("documents").insert(records).execute()
            inserted_count = len(result.data)
            logs.append(f"Successfully inserted {inserted_count} rows.")
            return inserted_count, logs
        return 0, logs

    def ingest_website(self, request: WebsiteIngestionRequest) -> IngestionResponse:
        logs = ["Starting website ingestion..."]
        try:
            text = self.scrape_website(request.url)
            logs.append(f"Scraped {len(text)} characters.")
            
            classification = self.classify_document(text, request.url)
            summary = self.generate_summary(text, request.url)
            
            cat = request.category or classification.get("category", "general")
            ind = request.industry or classification.get("industry", "general")
            title = classification.get("title", request.url)
            
            metadata = {
                "source": request.url,
                "type": "website",
                "url": request.url,
                "title": title,
                "category": cat,
                "industry": ind,
                "topics": classification.get("topics", []),
            }
            
            num_chunks, storage_logs = self.process_and_store(text, metadata)
            logs.extend(storage_logs)
            
            self.save_manifest_entry(
                source=request.url,
                source_type="website",
                title=title,
                category=cat,
                industry=ind,
                summary=summary,
                chunk_count=num_chunks
            )
            logs.append("Saved manifest entry.")
            return IngestionResponse(
                status="success",
                message="Website ingested successfully",
                chunk_count=num_chunks,
                logs=logs
            )
        except Exception as e:
            return IngestionResponse(
                status="error",
                message=str(e),
                chunk_count=0,
                logs=logs
            )

    def ingest_file(self, filename: str, file_bytes: bytes) -> IngestionResponse:
        logs = [f"Starting ingestion for {filename}"]
        ext = filename.split('.')[-1].lower()
        storage_path = None
        
        try:
            storage_path = self.upload_file_to_storage(file_bytes, filename)
            logs.append(f"Uploaded to storage: {storage_path}")
        except Exception as e:
            logs.append(f"Could not upload to storage: {str(e)}")

        try:
            if ext == 'pdf':
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                text = ""
                for page in doc:
                    text += page.get_text() + "\n"
                
                if len(text.strip()) < 100:
                    logs.append("PDF text is sparse, attempting Gemini Multimodal OCR fallback...")
                    from google import genai
                    from google.genai import types
                    client = genai.Client()
                    response = client.models.generate_content(
                        model='gemini-2.5-flash-lite',
                        contents=[
                            types.Part.from_bytes(data=file_bytes, mime_type='application/pdf'),
                            "OCR exact text extraction."
                        ]
                    )
                    text = response.text or ""
            elif ext == 'docx':
                doc = docx.Document(io.BytesIO(file_bytes))
                text = "\n".join([para.text for para in doc.paragraphs])
            elif ext == 'xlsx':
                df = pd.read_excel(io.BytesIO(file_bytes))
                text = df.to_string()
            elif ext == 'pptx':
                prs = Presentation(io.BytesIO(file_bytes))
                text_parts = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text:
                            text_parts.append(shape.text)
                text = "\n".join(text_parts)
            elif ext in ['txt', 'md']:
                text = file_bytes.decode('utf-8', errors='ignore')
            else:
                raise ValueError(f"Unsupported file format: {ext}")

            classification = self.classify_document(text, filename)
            summary = self.generate_summary(text, filename)
            
            metadata = {
                "source": filename,
                "type": ext,
                "title": classification.get("title", filename),
                "category": classification.get("category", "general"),
                "industry": classification.get("industry", "general"),
                "topics": classification.get("topics", []),
            }
            if storage_path:
                metadata["storage_path"] = storage_path
                
            num_chunks, storage_logs = self.process_and_store(text, metadata)
            logs.extend(storage_logs)
            
            self.save_manifest_entry(
                source=filename,
                source_type=ext,
                title=classification.get("title", filename),
                category=classification.get("category", "general"),
                industry=classification.get("industry", "general"),
                summary=summary,
                chunk_count=num_chunks
            )
            logs.append("Saved manifest entry.")
            return IngestionResponse(
                status="success",
                message=f"File {filename} ingested successfully",
                chunk_count=num_chunks,
                logs=logs
            )
        except Exception as e:
            return IngestionResponse(
                status="error",
                message=str(e),
                chunk_count=0,
                logs=logs
            )

    def get_ingested_documents(self) -> list[IngestedDocResponse]:
        all_data = []
        limit = 1000
        offset = 0
        while True:
            response = self.supabase.client.table("documents").select("id, content, metadata").range(offset, offset + limit - 1).execute()
            batch = response.data
            if not batch:
                break
            for row in batch:
                all_data.append(IngestedDocResponse(
                    id=str(row["id"]),
                    content=row["content"],
                    metadata=row["metadata"] or {}
                ))
            if len(batch) < limit:
                break
            offset += limit
        return all_data
